"""Cross-check saved deliverables against original processing records."""
from pathlib import Path
import json,xml.etree.ElementTree as ET
import numpy as np
import pandas as pd
import fitz
from PIL import Image
from openpyxl import load_workbook

OUT=Path(__file__).resolve().parents[1]

def main():
    df=pd.read_csv(OUT/'data/逐图计数与质量.csv')
    inv=pd.read_csv(OUT/'data/图片清单.csv')
    assert len(df)==26403 and df.frame_id.is_unique
    assert set(df.frame_id)==set(inv.frame_id)
    assert df.source_sha256.str.fullmatch('[0-9a-f]{64}').all()
    assert df.source_unchanged.all()
    assert df.loc[df.group.isin(['B','C']),['left_count','right_count']].isna().all().all()
    assert df.other_count.isna().all()
    samples={r['frame_id']:r for r in map(json.loads,(OUT/'data/run_4e84dd351c3bac11/samples_checkpoint.jsonl').open())}
    matched=0
    for line in (OUT/'data/run_24cc7e7f926b5b7e/full_checkpoint.jsonl').open():
        r=json.loads(line)
        if r['frame_id'] in samples:
            s=samples[r['frame_id']]
            for key in ['sha256','auto_left_count','auto_right_count','auto_other_count']:
                assert r[key]==s[key],(r['frame_id'],key)
            matched+=1
    assert matched==150
    stats={}
    for name in ['主统计_仅初检达标范围','候选统计_未通过验收']:
        folder=OUT/'data'/name
        bins=pd.read_csv(folder/'全部时间窗统计.csv')
        overall=pd.read_csv(folder/'总体汇总.csv')
        mapping=pd.read_csv(folder/'拼接时间与原始时间对应.csv')
        gaps=pd.read_csv(folder/'分析时间缺口.csv')
        for _,o in overall.iterrows():
            b=bins[bins.group==o.group];m=mapping[mapping.group==o.group]
            if o.effective_hours==0:
                assert b.empty and pd.isna(o.left_mean) and pd.isna(o.left_share)
                continue
            for (axis,step),a in b.groupby(['axis','window_seconds']):
                assert np.isclose(a.observed_seconds.sum()/3600,o.effective_hours)
                for side in ['left','right']:
                    assert np.isclose(a[f'{side}_insect_minutes'].sum(),o[f'{side}_insect_minutes'])
                shares=a.left_share+a.right_share
                assert np.allclose(shares.dropna(),1)
                no_observation=a.observed_seconds==0
                assert a.loc[no_observation,['left_mean','right_mean','left_insect_minutes','right_insect_minutes']].isna().all().all()
                assert (a.observed_seconds<=a.end_s-a.start_s+1e-6).all()
            assert np.allclose(m.actual_end_s-m.actual_start_s,m.effective_end_s-m.effective_start_s)
            assert np.isclose(m.effective_start_s.iloc[0],0)
            assert np.allclose(m.effective_start_s.iloc[1:],m.effective_end_s.iloc[:-1])
            assert np.isclose(m.effective_end_s.iloc[-1]/3600,o.effective_hours)
            assert np.isclose(gaps[gaps.group==o.group].seconds.sum()/3600+o.effective_hours,o.actual_hours)
        stats[name]={'windows':len(bins),'segments':len(mapping),'integral_conservation':True,'no_gap_integration':True}
    figures=[]
    for p in sorted((OUT/'figures').glob('*.png')):
        im=Image.open(p);dpi=im.info.get('dpi',(0,0))
        assert all(abs(x-300)<.1 for x in dpi),(p,dpi)
        svg=p.with_suffix('.svg');pdf=p.with_suffix('.pdf')
        assert svg.exists() and pdf.exists() and pdf.stat().st_size>1000
        root=ET.parse(svg).getroot();texts=root.findall('.//{http://www.w3.org/2000/svg}text')
        assert texts,'SVG must retain editable text'
        with fitz.open(pdf) as doc:
            assert len(doc)==1
            pdf_text=doc[0].get_text()
            assert '褐飞虱' in pdf_text and '时间' in pdf_text
        if '有效时间拼接' in p.name:
            assert '衔接' in ''.join(root.itertext())
            assert '衔接' in pdf_text
        if '待复核' in p.name:
            assert '待复核' in ''.join(root.itertext())
        if '01_时间曲线' in p.name:
            for color in ['#2677a8','#d86c32']:
                data_paths=[e for e in root.findall('.//{http://www.w3.org/2000/svg}path')
                            if f'stroke: {color}' in e.get('style','') and e.get('d','').count('L')>3]
                assert len(data_paths)==(3 if '待复核' in p.name else 1)
                assert all(e.get('d','').count('M')==1 for e in data_paths),'Effective curves must remain connected'
        figures.append({'file':p.name,'pixels':list(im.size),'dpi':list(dpi),'editable_svg_text':True,'pdf_chinese_text_extractable':True})
    assert len(figures)==6
    wb=load_workbook(OUT/'褐飞虱逐图计数与时间统计_含待复核结果.xlsx',read_only=True,data_only=True)
    ws=wb['逐图计数'];rows=ws.iter_rows(values_only=True);headers=list(next(rows))
    columns=['照片标识','左侧自动候选数','右侧自动候选数','左侧主计数','右侧主计数']
    pos=[headers.index(c) for c in columns]
    for values,(_,r) in zip(rows,df.iterrows(),strict=True):
        actual=[values[k] for k in pos]
        expected=[r.frame_id,r.auto_left_count,r.auto_right_count,r.left_count,r.right_count]
        for x,y in zip(actual,expected):
            assert (x is None and pd.isna(y)) or x==y
    wb.close()
    result={'all_frame_records':26403,'cached_samples_equal_native_full_frames':matched,
            'statistics':stats,'figures':figures,'excel_all_frame_counts_equal_csv':True,
            'count_acceptance':'A模型目视初检达标；B/C未通过；独立人工专家验收未完成'}
    (OUT/'qa/交付交叉核验.json').write_text(json.dumps(result,ensure_ascii=False,indent=2))
    print(json.dumps(result,ensure_ascii=False,indent=2))

if __name__=='__main__':main()

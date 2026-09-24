"""Chinese deliverables with explicit measurement and review boundaries."""
from pathlib import Path
import json,hashlib
import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Font,PatternFill,Alignment
from openpyxl.utils import get_column_letter

OUT=Path(__file__).resolve().parents[1]
CN={
'frame_id':'照片标识','group':'实验组编号','group_label':'实验组','sequence_index':'组内序号_从0起',
'source_path':'原图路径','relative_path':'相对路径','size_bytes':'文件字节数','mtime_ns':'源文件修改时间_ns',
'width':'像素宽','height':'像素高','orientation':'EXIF方向','timestamp':'相机拍摄时间',
'metadata_error':'元数据异常','interval_s':'距前帧秒数','elapsed_s':'真实经过秒数','timestamp_duplicate':'拍摄时间重复',
'sample_id':'样本编号','processing_status':'处理状态','processed_at':'处理时间','source_sha256':'源文件SHA256',
'source_unchanged':'处理时大小及修改时间与清单一致','detector_version':'检测版本','detector_fingerprint':'检测配置指纹',
'quality_ok':'通过图片质量初筛','quality_reason':'图片质量说明','brightness':'亮度指标','sharpness':'清晰度指标','scene_overlap':'植株区域匹配比例',
'auto_left_count':'左侧自动候选数','auto_right_count':'右侧自动候选数','auto_other_count':'其他位置自动候选数',
'reflection_or_bottom_candidates':'水面以下或倒影候选数','suspect_count':'需要复核的检测目标数',
'left_count':'左侧主计数','right_count':'右侧主计数','other_count':'其他位置确认数',
'other_missing_reason':'其他位置确认数留空原因','unannotated_preview':'未标注预览图路径',
'valid_primary':'纳入主统计','valid_candidate_summary':'纳入待复核候选统计','primary_missing_reason':'主计数留空原因',
'review_status':'复核状态','reference_left_count':'模型目视参考_左','reference_right_count':'模型目视参考_右',
'reference_notes':'模型目视说明','expert_reviewed':'实验人员已复核','review_required':'需要复核',
'annotation_path':'自动标注图路径','error':'处理错误','image_effective_elapsed_s':'图片质量合格累计秒数',
'primary_effective_elapsed_s':'主统计累计有效秒数','candidate_effective_elapsed_s':'候选统计累计有效秒数',
'axis':'时间轴','window_seconds':'时间窗秒数','start_s':'窗口起点秒','end_s':'窗口终点秒',
'start_hour':'窗口起点小时','end_hour':'窗口终点小时','observed_seconds':'实际纳入秒数','coverage':'观察覆盖比例',
'interval_pieces':'积分小段数','left_mean':'左侧时间加权均值_只','right_mean':'右侧时间加权均值_只',
'left_insect_minutes':'左侧累计占据量_虫分钟','right_insect_minutes':'右侧累计占据量_虫分钟',
'left_share':'左侧占据量占比_0到1','right_share':'右侧占据量占比_0到1','data_status':'数据状态',
'n_images':'照片总数','n_valid':'可纳入帧数','n_invalid':'未纳入帧数','first_timestamp':'首帧相机时间','last_timestamp':'末帧相机时间',
'actual_hours':'真实经过小时','effective_hours':'纳入统计有效小时','segment':'连续片段编号',
'actual_start_s':'真实时间起点秒','actual_end_s':'真实时间终点秒','effective_start_s':'有效时间起点秒','effective_end_s':'有效时间终点秒',
'first_frame':'片段首帧','last_frame':'片段末帧','before':'缺口前帧','after':'缺口后帧','seconds':'缺口秒数','reason':'原因',
'version':'检测版本','side':'侧别','reference_count':'模型目视参考数','auto_count':'自动计数','reference_status':'目视参考状态',
'notes':'说明','reviewer':'参考计数来源','inspection_before_predictions':'目视计数先于查看预测','independent_human_expert':'独立人工专家核验',
'selected_samples':'预选留出照片数','usable_reference_sides':'可用参考侧数','unresolved_reference_sides':'无法确定参考侧数',
'mae':'平均绝对误差_只','within_two_fraction':'误差不超过2只的比例',
'numerical_target_on_usable_sides':'可用侧数值达到目标','initial_model_visual_check_passed':'30张初始模型目视核验通过',
'evaluation_role':'核验用途'}

def text_number(value,decimals=2):
    return '未形成可靠统计' if pd.isna(value) else f'{value:.{decimals}f}'

def deliver(inventory,df,audit,comparison,metrics,primary,candidate):
    _,pb,pm,po,pg=primary;_,cb,cm,co,cg=candidate
    instructions=pd.DataFrame([
        ['结果状态','全量自动逐图处理完成；BPH33与回补组计数验收未通过，不能作为最终定量结果。'],
        ['左右分组','左侧对照、右侧抗性；A/B各侧5株，C各侧4株；仅报告侧合计。'],
        ['主计数','仅模型目视初始误差目标通过的A组自动计数；不是逐张人工确认。B/C主计数空白。'],
        ['自动候选','算法独立处理每张原图的候选数；可含漏计、误计及重叠问题。B/C图表均为待复核结果。'],
        ['空白','无法确认或未通过验收；不是0。其他位置确认数未经验收，统一留空，自动候选另列。'],
        ['疑似目标数','这是自动候选中的复核标记数，与左右或其他候选存在重叠，不能相加得到总虫数。'],
        ['有效时间与衔接','仅相邻合格帧0<间隔≤15秒积分。停拍时长去除；主曲线连接；横轴标“衔接”。'],
        ['五分钟均值','相邻帧计数作线性连接，用梯形积分分配到5分钟有效时间窗，再除以纳入秒数。'],
        ['小时占比','每侧虫数×秒的积分/两侧积分合计。表内0~1，图内百分比。分母为0时留空。'],
        ['核验范围','每组20张校准；另30张未用于初始训练的照片由Codex目视检查。并非人工专家盲法金标准。'],
        ['回补组清晰度','清晰度<90的帧从候选时间统计排除；这是复核触发阈值，尚未完成新的独立验证。'],
        ['实验重复','每个日期文件夹视为一次连续观察；连续照片不是独立重复，不做显著性检验。'],
        ['统计含义','株上可见数量与占据分布，不是吸食量、吸食速率或总虫群精确数量。'],
        ['坐标','全部自动检测坐标.csv.gz对应自动候选计数；参考计数另列，不假装已经人工修正全部坐标。'],
        ['原始数据','原图只读保留。相机EXIF时间未作时钟准确性或时区校准。'],
    ],columns=['项目','说明'])
    workbook=OUT/'褐飞虱逐图计数与时间统计_含待复核结果.xlsx'
    sheets={'使用说明':instructions,'图片清单':inventory,'逐图计数':df,
            '校准与留出照片':pd.read_csv(OUT/'data/校准与留出样本.csv'),
            '主统计_总体':po,'主统计_5分钟':pb[(pb.axis=='effective')&(pb.window_seconds==300)],
            '主统计_逐小时':pb[(pb.axis=='effective')&(pb.window_seconds==3600)],
            '待复核_总体':co,'待复核_5分钟':cb[(cb.axis=='effective')&(cb.window_seconds==300)],
            '待复核_逐小时':cb[(cb.axis=='effective')&(cb.window_seconds==3600)],
            '待复核_真实时间5分钟':cb[(cb.axis=='actual')&(cb.window_seconds==300)],
            '拍摄停拍缺口':pd.read_csv(OUT/'data/拍摄时间缺口.csv'),
            '候选统计排除间隔':cg,'衔接时间对应':cm,
            '核验汇总':metrics,'逐样本核验':comparison.rename(columns={'error':'计数差_自动减参考'}),
            '标注图索引':pd.read_csv(OUT/'data/标注样本索引.csv'),
            '优先复核照片':pd.read_csv(OUT/'data/优先复核照片.csv'),
            '异常与复核':pd.read_csv(OUT/'data/异常与复核清单.csv'),
            '字段字典':pd.DataFrame(list(CN.items()),columns=['CSV字段','中文释义'])}
    with pd.ExcelWriter(workbook,engine='openpyxl') as writer:
        for name,t in sheets.items():
            t=t.copy()
            # Excel only preserves 15 significant digits in numeric cells.
            if 'mtime_ns' in t:t['mtime_ns']=t.mtime_ns.map(lambda v:str(int(v)) if pd.notna(v) else '')
            if 'axis' in t:t['axis']=t.axis.replace({'effective':'累计有效时间','actual':'真实经过时间'})
            if 'side' in t:t['side']=t.side.replace({'left':'左侧','right':'右侧'})
            if 'reason' in t:t['reason']=t.reason.replace({'acquisition_gap':'停拍缺口','quality_gap':'计数验收或画面质量不满足纳入条件'})
            t.rename(columns=CN).to_excel(writer,sheet_name=name,index=False)
        for ws in writer.book:
            ws.freeze_panes='D2' if ws.title in ['图片清单','逐图计数','异常与复核'] else 'A2'
            ws.auto_filter.ref=ws.dimensions
            ws.sheet_view.zoomScale=85
            for cell in ws[1]:cell.font=Font(name='Arial',bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='31566D');cell.alignment=Alignment(wrap_text=True,vertical='center')
            ws.row_dimensions[1].height=42
            for j,col in enumerate(ws.iter_cols(min_row=1,max_row=min(ws.max_row,35)),1):
                head=str(col[0].value or '')
                width=50 if any(x in head for x in ['路径','原因','说明','状态']) else 24 if '时间' in head else min(26,max(14,len(head)*1.4))
                ws.column_dimensions[get_column_letter(j)].width=width
            if ws.title=='使用说明':
                ws.column_dimensions['A'].width=24;ws.column_dimensions['B'].width=115
                for row in ws.iter_rows(min_row=2):
                    row[1].alignment=Alignment(wrap_text=True,vertical='center');ws.row_dimensions[row[0].row].height=34
    # Read the saved workbook, not the in-memory object, for row/count QA.
    wb=load_workbook(workbook,read_only=True,data_only=True)
    assert wb['逐图计数'].max_row==26404 and wb['图片清单'].max_row==26404
    assert wb['核验汇总'].max_row==len(metrics)+1
    wb.close()
    a=po[po.group=='A'].iloc[0]
    metadata=json.loads((OUT/'data/inventory_summary.json').read_text())
    timeline='\n'.join(f'| {r["group"]} | {r["n"]:,} | {r["elapsed_hours"]:.3f} | {r["potentially_observed_hours"]:.3f} |' for r in metadata)
    validations=[]
    for g in 'ABC':
        l=metrics[(metrics.group==g)&(metrics.side=='left')&(metrics.version=='v1')].iloc[0]
        r=metrics[(metrics.group==g)&(metrics.side=='right')&(metrics.version=='v1')].iloc[0]
        validations.append(f'| {g} | {int(l.usable_reference_sides)}/30 | {l.mae:.2f} | {l.within_two_fraction:.1%} | {int(r.usable_reference_sides)}/30 | {r.mae:.2f} | {r.within_two_fraction:.1%} | '+('初始模型自检达标' if g=='A' else '未通过')+' |')
    report=f'''# 褐飞虱逐图计数与时间分布分析

生成日期：2026-09-16。状态：**全量自动检测完成，整项计数验收尚未完成**。

26,403张原图均有独立处理记录、原图路径、相机时间和自动检测坐标。ZH11 B3-OE组达到初始模型目视核验误差目标；BPH33及BPH33回补组未通过，因此其主计数与主统计留空。三组连续曲线与热图另以“自动候选计数，待复核”版本提供，不能当作已验收的科研结果。

## 1. 数据范围

| 组别 | 照片数 | 真实经过小时 | 仅按拍摄间隔可用小时 |
|---|---:|---:|---:|
{timeline}

A=ZH11 B3-OE，B=BPH33，C=BPH33回补。A/B各侧5株，C各侧4株。全部按左侧对照、右侧抗性，仅输出侧合计。各目录视为一场连续观察；目录名日期与首帧相机日期可能不同，以EXIF排序。相机时钟准确性、时区未另行校准。未将连续照片作为独立实验重复。

## 2. 逐图检测与追溯

在本地读取4416×2488原始像素，每张照片独立进行颜色、局部对比度、植株空间区域及虫体形态检测。A采用校准后的连通域与分割；B/C增加本地图像训练的HOG线性分类器。没有用抽样计数、邻帧复制或时间插值填满逐图虫数。背景和植株区域由每组20张校准图建立；150个清楚单体训练点来自校准图，训练精度不作为计数验收精度。

保留了v1完整参数、权重、背景、代码指纹和断点记录。v2尝试保留静止虫体的原始局部对比度，仅完成150张样本对照，仍未通过B/C计数验收，未替换全量v1结果。低分位背景仍可能降低始终静止且与背景混合的虫体信号，这是未消除的局限，不能声称静止虫体已全部找回。

株上候选按左右分类；底部、水面以下与疑似倒影另列。其他位置候选不是完整虫数金标准，确认数留空。自动框、疑似标记和原图路径保留，重叠、截断、植株边缘和基部候选进入复核。疑似目标数是候选中的复核标记数，与左右或其他位置候选有重叠，不能相加当作总虫数。某个候选被标出，并不等于已经确认是一个独立褐飞虱。

## 3. 计数核验

每组先预选20张校准图，另外30张作初始留出核验。校准图均用于画面/背景校准，其中部分用于单体形态训练及逐体校准；不能称60张校准图均已逐体精确人工标注。留出图由Codex先看未标注全图，必要时看原始像素裁剪，记录参考计数后再比较预测。**这是模型目视自检，不是独立实验人员盲法核验。**

| 组 | 左可用参考侧 | 左MAE | 左误差≤2比例 | 右可用参考侧 | 右MAE | 右误差≤2比例 | 结论 |
|---|---:|---:|---:|---:|---:|---:|---|
{chr(10).join(validations)}

初始目标为每侧MAE≤1只且至少95%的误差≤2只，并要求预定30张都能形成参考计数。B/C不少画面有密集叠靠或归属不清，参考数留空，没有把估计整数或零当真值。可用侧的误差只是部分参考，不能代表未能分辨的画面。回补组右侧可用样本数值较好，仍不满足全组30张两侧完整验收。

失败后尝试原始对比度修正并检查B组低分候选。降低阈值可找回一些虫体，也引入叶尖和箱体条纹；没有据此把同一留出样本重复调参后的成绩称为新的独立验收。**B/C仍需清楚逐体标注、密集重叠处理和新的留出样本验收。** 全部照片尚未由实验人员逐图复核；本交付没有宣称完成这一步。

## 4. 时间处理与“衔接”

只对相邻两张可纳入照片且0<间隔≤15秒的区间作梯形积分。停拍、失焦和无效区间不赋值、不跨越积分。主图横轴使用累计有效观察时间，压缩掉这些间隔；曲线连续连接，横轴用短刻线和“衔接”注明拼接位置。连线仅帮助阅读，不给缺口创造观察时长或虫数。真实时间附图保留间断，并按连续片段分别绘线。

每5分钟：把虫数×秒积分精确分配到有效时间窗，再除以实际纳入秒数。每小时：左/右积分分别除以两侧积分总和。两侧均为零时，占比留空；没有观察的时段，各统计量留空。保留末尾不足一小时的窗口，没有统一截取24小时。拼接前后对应、片段首末帧和排除间隔均保存。

候选图使用统一的图片质量规则；C清晰度指标<90的帧暂不纳入，因为核验中发现明显失焦。此规则只作复核触发，尚未完成新的独立准确性验证。密集重叠问题不因画面整体清楚而消失，所以B/C候选图仍标“待复核”。

回补组共有{int(co.loc[co.group=="C","n_invalid"].iloc[0])}张未通过图片质量初筛；排除这些帧及其相邻不可积分区间后，候选曲线纳入{co.loc[co.group=="C","effective_hours"].iloc[0]:.3f}小时。密集短刻线表示多次剔除无效间隔，完整位置保留在时间对应表中。

## 5. 当前可报告的范围

A组在初始模型自检口径下纳入{a.effective_hours:.3f}小时，左侧时间加权均值{text_number(a.left_mean)}只、右侧{text_number(a.right_mean)}只；累计占据量左侧{a.left_share:.1%}、右侧{a.right_share:.1%}。这些是本场观察的可见虫数和株上占据分布，仍需实验人员抽查与确认。

B/C未形成通过验收的主统计。可查看候选曲线定位变化和复核时段，但本报告不据候选数给出抗性效应大小或显著性结论。三组株数和画面不同，不能直接把侧合计当成可跨实验无条件比较的统一指标。照片中的停留不等同于实际吸食量或吸食速率。

## 6. 文件与复现

- `褐飞虱逐图计数与时间统计_含待复核结果.xlsx`：中文字段、逐图数据、主统计、候选统计、核验与复核清单。
- `data/逐图计数与质量.csv`：完整26,403行，自动候选数和主计数分列，缺失为空。
- `data/主统计_仅初检达标范围/`：通过初始模型目视核验范围的统计。
- `data/候选统计_未通过验收/`：全量候选的探索性时间统计。
- `figures/`：两套时间曲线、逐小时热图、真实时间附图；SVG/PDF/300 dpi PNG。
- `data/全部自动检测坐标.csv.gz`：原图像素坐标。完整候选、拒绝候选和理由在全量JSONL中。
- `evidence/`：校准/留出样本自动标注图。`qa/`保留未标注参考视图、局部放大和验证记录。
- `scripts/`、`cache/`、`data/project_config.json`：复现脚本、局部模型和参数。原图始终位于原目录。

复现命令见“复现说明.md”。本机照片复核工具使用方法见“review/复核工具使用说明.md”。重新检测支持断点续跑；坐标与全部自动计数一致。源文件大小和修改时间在处理时与初始清单比对，读取内容保存SHA256；未改写原图。总处理错误{audit['processing_errors']}，缺失记录{audit['missing_records']}，自动坐标{audit['coordinate_count']:,}个。技术记录完整不等于B/C计数精度已合格。
'''
    (OUT/'中文分析报告.md').write_text(report)
    (OUT/'README_先读.md').write_text('''# 结果状态与入口

全部26,403张已自动逐图处理。**BPH33及回补组未通过计数验收，本交付包含待复核结果，不是全部完成的最终科研计数。**

先看“中文分析报告.md”和Excel“使用说明”。主计数和自动候选数分列，无法确认的主计数留空。

主图按有效时间衔接，横轴有“衔接”标识。带“待复核”的三组图是算法候选结果，不能用于正式定量结论。未带此后缀的主统计图仅显示初始模型自检达标的A组，其余组留空。

原图未修改。Excel、CSV、可编辑SVG、矢量PDF和300 dpi PNG均在此目录或子目录。详细复核材料、模型目视核验范围和未完成事项见报告。
''')
    (OUT/'复现说明.md').write_text('''# 本地复现

当前已安装环境：`/opt/miniconda3/bin/python`。无需外传图片。必须挂载原始照片所在卷。

在本结果目录执行：

```sh
/opt/miniconda3/bin/python scripts/run_batch.py --mode full --workers 16
/opt/miniconda3/bin/python scripts/test_statistics.py
/opt/miniconda3/bin/python scripts/export_results.py
```

第一步按源图逐张处理并断点续跑；第二步仅运行合成数组的数值检查，不生成实验数据；第三步要求全量记录齐全后导出。检测器指纹由检测脚本、权重和背景文件的SHA256组成。参数见全量运行目录内detector_manifest.json。

不能通过改成0或复制邻帧来填补主计数。若修改检测参数或用现有留出图调参，必须建立新版本并用新留出样本核验；不能复用旧版通过结论。若增加人工复核，应保存复核人、时间、原图、逐体坐标和修正前后计数，再更新统计；目前没有宣称已经完成独立人工复核。

完整JSONL包含被拒绝的候选；坐标CSV只导出最终自动候选。`v2`仅为样本级探索性尝试，未替换全量`v1`。
''')
    audit['xlsx_reopened']=True;audit['xlsx_frame_rows']=26403
    (OUT/'qa/全量记录验收.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2))
    manifest=[]
    for path in sorted(OUT.rglob('*')):
        if path.is_file() and path.suffix in ['.py','.json','.csv','.xlsx','.md','.svg','.pdf','.png'] and 'cache/samples' not in str(path) and '__pycache__' not in str(path):
            if path.name=='交付文件清单.json':continue
            manifest.append(dict(path=str(path.relative_to(OUT)),bytes=path.stat().st_size,sha256=hashlib.sha256(path.read_bytes()).hexdigest()))
    (OUT/'qa/交付文件清单.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2))

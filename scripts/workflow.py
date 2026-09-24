#!/usr/bin/env python3
"""Local, source-preserving replay of the frozen 2023 BPH analysis.

The frozen detector is scene-specific. New-study scaffolds cannot run it.
No network, dependency installation, or source-image writes are performed.
"""
from __future__ import annotations
import argparse
import csv
import gzip
import hashlib
import importlib.metadata
import importlib.util
import json
import os
from pathlib import Path
import shutil
import subprocess
import sys
from datetime import datetime

SKILL = Path(__file__).resolve().parents[1]
FROZEN = SKILL / 'assets/frozen_v1'
CASE = SKILL / 'assets/case_2023'
PROFILE = 'fafu-2023-v1'
FINGERPRINT = '24cc7e7f926b5b7e'
TOTAL = 26403
MARKER = '.bph_skill_project.json'
IMAGE_EXTENSIONS = {'.jpg', '.jpeg', '.png', '.tif', '.tiff', '.bmp', '.webp'}


def now():
    return datetime.now().astimezone().isoformat()


def digest(path):
    h = hashlib.sha256()
    with Path(path).open('rb') as f:
        for block in iter(lambda: f.read(1024 * 1024), b''):
            h.update(block)
    return h.hexdigest()


def write_json(path, value):
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(value, ensure_ascii=False, indent=2), encoding='utf-8')


def read_json(path):
    return json.loads(Path(path).read_text(encoding='utf-8'))


def verify_assets():
    manifest = read_json(SKILL / 'assets/asset_manifest.json')
    for row in manifest['files']:
        p = SKILL / row['path']
        if not p.is_file() or digest(p) != row['sha256']:
            raise ValueError(f'冻结资源缺失或已改变：{p}。不可沿用旧验收结论。')
    return len(manifest['files'])


def doctor():
    modules = {'numpy': 'numpy', 'pandas': 'pandas', 'cv2': 'opencv-python',
               'scipy': 'scipy', 'skimage': 'scikit-image', 'sklearn': 'scikit-learn',
               'matplotlib': 'matplotlib', 'openpyxl': 'openpyxl', 'PIL': 'Pillow',
               'joblib': 'joblib', 'fitz': 'PyMuPDF'}
    result = {'python': sys.version.split()[0], 'executable': sys.executable,
              'modules': {}, 'missing': [], 'asset_files_verified': verify_assets()}
    for module, package in modules.items():
        if importlib.util.find_spec(module) is None:
            result['missing'].append(module)
            continue
        try:
            result['modules'][module] = importlib.metadata.version(package)
        except importlib.metadata.PackageNotFoundError:
            result['modules'][module] = 'installed; distribution name differs'
    if 'matplotlib' not in result['missing']:
        from matplotlib import font_manager
        try:
            result['chinese_font'] = font_manager.findfont('Arial Unicode MS', fallback_to_default=False)
        except ValueError:
            result['chinese_font'] = None
    result['reference_environment'] = read_json(CASE / 'data/运行环境版本.json')
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if result['missing']:
        raise ValueError('当前环境缺少依赖；本工具不会自动安装。先检查现有 Python 环境。')
    if not result.get('chinese_font'):
        raise ValueError('冻结绘图配置需要可用的 Arial Unicode MS；请使用已有中文字体环境或另建绘图配置。')
    return result


def create_output(output, mode, source_roots=()):
    output = Path(output).expanduser().resolve()
    if output == SKILL or output.is_relative_to(SKILL):
        raise ValueError('输出不能写入 Skill 自身。')
    for root in source_roots:
        root = Path(root).expanduser().resolve()
        if output == root or output.is_relative_to(root):
            raise ValueError('输出不能写入原始照片目录。')
    marker = output / MARKER
    if output.exists() and any(output.iterdir()):
        if not marker.exists() or read_json(marker).get('mode') != mode:
            raise ValueError(f'拒绝覆盖已有目录：{output}。请选择新的输出目录。')
    else:
        output.mkdir(parents=True, exist_ok=True)
        write_json(marker, {'mode': mode, 'created_at': now(), 'skill_version': '1.0.0'})
    for name in ['data', 'qa', 'figures', 'evidence', 'cache', 'review']:
        (output / name).mkdir(exist_ok=True)
    return output


def project(output, mode='raw-replay'):
    output = Path(output).expanduser().resolve()
    info = read_json(output / MARKER)
    if info.get('mode') != mode:
        raise ValueError(f'该命令仅支持 {mode} 项目，不能给新照片套用历史模型。')
    return output, info


def copy_runtime(output, models=False):
    shutil.copytree(FROZEN / 'scripts', output / 'scripts', dirs_exist_ok=True,
                    ignore=shutil.ignore_patterns('__pycache__', '*.pyc'))
    if models:
        shutil.copytree(FROZEN / 'cache', output / 'cache', dirs_exist_ok=True)
        shutil.copytree(FROZEN / 'review', output / 'review', dirs_exist_ok=True)


def copy_case_data(output):
    for p in (CASE / 'data').rglob('*'):
        if not p.is_file():
            continue
        rel = p.relative_to(CASE / 'data')
        target = output / 'data' / rel
        target.parent.mkdir(parents=True, exist_ok=True)
        if p.suffix == '.gz':
            with gzip.open(p, 'rb') as inf, target.with_suffix('').open('wb') as outf:
                shutil.copyfileobj(inf, outf)
        else:
            shutil.copy2(p, target)


def run_script(output, filename, *args):
    command = [sys.executable, '-u', str(output / 'scripts' / filename), *map(str, args)]
    subprocess.run(command, cwd=output, check=True)


def baseline():
    import pandas as pd
    return pd.read_csv(CASE / 'data/逐图计数与质量.csv.gz')


def compare_statistics(output):
    import pandas as pd
    checked = []
    for kind, name in [('primary', '主统计_仅初检达标范围'), ('candidate', '候选统计_未通过验收')]:
        for ref in sorted((CASE / 'expected_statistics' / kind).glob('*.csv.gz')):
            actual = output / 'data' / name / ref.name.removesuffix('.gz')
            a, b = pd.read_csv(actual), pd.read_csv(ref)
            pd.testing.assert_frame_equal(a, b, check_dtype=False, check_exact=False, rtol=1e-11, atol=1e-8)
            checked.append({'table': str(actual.relative_to(output)), 'rows': len(a)})
    write_json(output / 'qa/统计复现核对.json', {
        'checked_at': now(), 'matched_tables': checked,
        'meaning': '与冻结逐图数据的历史统计一致；不代表B/C计数准确性通过验收'})
    return checked


def check_figures(output):
    from PIL import Image
    import fitz
    import xml.etree.ElementTree as ET
    checked = []
    ns = {'s': 'http://www.w3.org/2000/svg'}
    for path in sorted((output / 'figures').glob('*.png')):
        with Image.open(path) as im:
            assert all(abs(v - 300) < .1 for v in im.info['dpi'])
        root = ET.parse(path.with_suffix('.svg')).getroot()
        assert root.findall('.//s:text', ns)
        text = ''.join(root.itertext())
        with fitz.open(path.with_suffix('.pdf')) as doc:
            assert len(doc) == 1 and '褐飞虱' in doc[0].get_text()
        if '有效时间拼接' in path.name:
            assert '衔接' in text
        if '待复核' in path.name:
            assert '待复核' in text
        if path.name.startswith('01_'):
            for color in ['#2677a8', '#d86c32']:
                lines = [p for p in root.findall('.//s:path', ns)
                         if f'stroke: {color}' in p.get('style', '') and p.get('d', '').count('L') > 3]
                assert len(lines) == (3 if '待复核' in path.name else 1)
                assert all(p.get('d', '').count('M') == 1 for p in lines)
        checked.append(path.name)
    assert len(checked) == 6
    write_json(output / 'qa/图形技术检查.json', {'files': checked, 'formats': ['SVG', 'PDF', '300dpi PNG'],
               'visual_inspection_still_required': True})


def check_replot_workbook(output, df):
    """Check saved Excel values, including missing counts, against the CSV."""
    import pandas as pd
    from openpyxl import load_workbook
    from report_results import CN
    wb = load_workbook(output / '历史计数复绘_含待复核结果.xlsx', read_only=True, data_only=True)
    try:
        ws = wb['历史逐图计数']
        assert ws.max_row == len(df) + 1
        rows = ws.iter_rows(values_only=True)
        headers = list(next(rows))
        keys = ['frame_id', 'auto_left_count', 'auto_right_count', 'left_count', 'right_count']
        positions = [headers.index(CN.get(key, key)) for key in keys]
        for values, expected in zip(rows, df[keys].itertuples(index=False, name=None), strict=True):
            for position, value in zip(positions, expected, strict=True):
                actual = values[position]
                assert (actual is None and pd.isna(value)) or actual == value
    finally:
        wb.close()
    write_json(output / 'qa/复绘Excel核对.json', {
        'rows_checked': len(df), 'fields_checked': keys, 'matches_csv': True})


def refresh_delivery_manifest(output):
    """Hash final report and checks after the replay wrapper adds provenance."""
    rows = []
    for path in sorted(output.rglob('*')):
        if (path.is_file() and path.suffix in {'.py', '.json', '.csv', '.xlsx', '.md', '.svg', '.pdf', '.png'}
                and '__pycache__' not in path.parts and 'cache/samples' not in str(path)
                and path.name != '交付文件清单.json'):
            rows.append({'path': str(path.relative_to(output)), 'bytes': path.stat().st_size, 'sha256': digest(path)})
    write_json(output / 'qa/交付文件清单.json', rows)


def replot(output):
    verify_assets()
    output = create_output(output, 'statistics-replay')
    copy_runtime(output)
    copy_case_data(output)
    sys.path.insert(0, str(output / 'scripts'))
    import pandas as pd
    import export_results as exporter
    from report_results import CN
    df = pd.read_csv(output / 'data/逐图计数与质量.csv')
    df['timestamp'] = pd.to_datetime(df.timestamp)
    assert len(df) == TOTAL and df.frame_id.is_unique
    primary = exporter.tables_and_figures(df)
    candidate = exporter.tables_and_figures(df, candidate=True)
    checked = compare_statistics(output)
    check_figures(output)
    with pd.ExcelWriter(output / '历史计数复绘_含待复核结果.xlsx', engine='openpyxl') as writer:
        pd.DataFrame([['模式', '仅用已存逐图数据复绘，没有重新检测原图或重新进行目视核验'],
                      ['计数状态', 'A仅历史模型目视初检达标；B/C未通过；没有独立人工专家验收'],
                      ['时间', '缺口不积分；有效时间曲线连接，横轴标衔接；真实时间附图保留缺口'],
                      ['来源', 'Skill中冻结的2026-09-16逐图计数快照，原图未打包']],
                     columns=['项目', '说明']).to_excel(writer, sheet_name='先读说明', index=False)
        table = df.copy()
        table['mtime_ns'] = table.mtime_ns.astype(str)
        table.rename(columns=CN).to_excel(writer, sheet_name='历史逐图计数', index=False)
        for prefix, result in [('主统计', primary), ('候选_待复核', candidate)]:
            _, bins, mapping, overall, gaps = result
            overall.rename(columns=CN).to_excel(writer, sheet_name=prefix + '_总体', index=False)
            for step, label in [(300, '5分钟'), (3600, '逐小时')]:
                bins[(bins.axis == 'effective') & (bins.window_seconds == step)].rename(columns=CN).to_excel(
                    writer, sheet_name=prefix + '_' + label, index=False)
            mapping.rename(columns=CN).to_excel(writer, sheet_name=prefix + '_衔接对应', index=False)
        pd.read_csv(output / 'data/计数核验汇总.csv').rename(columns=CN).to_excel(writer, sheet_name='历史核验状态', index=False)
        for ws in writer.book:
            ws.freeze_panes = 'A2'; ws.auto_filter.ref = ws.dimensions
    check_replot_workbook(output, df)
    (output / '复绘说明.md').write_text(
        '# 历史逐图计数复绘\n\n本次仅复绘冻结的26,403行计数，没有重新读取原始照片或重新进行目视核验。'
        '三组候选图保留待复核标识；A仅历史模型目视初检达标，B/C未通过。'
        '原始照片及完整逐体坐标不属于此复绘输出；原图重跑请使用 prepare / run / export。\n\n'
        f'运行时间：{now()}。{len(checked)}张统计表与冻结基准一致。技术检查不能替代计数准确性验收。\n', encoding='utf-8')
    refresh_delivery_manifest(output)
    print('统计复绘完成：', output)


def prepare(config_path, output):
    verify_assets()
    config = read_json(config_path)
    if config.get('profile') != PROFILE or set(config.get('source_folders', {})) != set('ABC'):
        raise ValueError('原图复现只接受 fafu-2023-v1 的A/B/C三组；新实验请用 init-new。')
    roots = {g: Path(p).expanduser().resolve() for g, p in config['source_folders'].items()}
    if any(not p.is_dir() for p in roots.values()):
        raise ValueError('原始照片卷或目录不可用；请挂载或修改配置中的三个目录。')
    output = create_output(output, 'raw-replay', roots.values())
    info = read_json(output / MARKER)
    normalized = {g: str(p) for g, p in roots.items()}
    if info.get('source_folders') not in [None, normalized]:
        raise ValueError('已有项目的输入目录不同；请新建输出目录。')
    if info.get('prepared'):
        check_runtime(output)
        print('复现项目已准备，可继续 smoke / run / export：', output)
        return output
    copy_runtime(output, models=True)
    copy_case_data(output)
    # The reference count snapshot must not masquerade as newly processed data.
    snapshot = output / 'data/逐图计数与质量.csv'
    snapshot.rename(output / 'data/历史逐图计数参考.csv')
    import pandas as pd
    inv = pd.read_csv(CASE / 'data/图片清单.csv.gz')
    rows = []
    for g, a in inv.groupby('group'):
        expected = {str(Path(p).relative_to(Path(p).parts[0])) for p in a.relative_path}
        actual = {str(p.relative_to(roots[g])) for p in roots[g].rglob('*')
                  if p.is_file() and p.suffix.lower() in IMAGE_EXTENSIONS and not p.name.startswith('._')}
        if actual != expected:
            raise ValueError(f'{g}组照片清单与冻结案例不一致；缺少{len(expected-actual)}张，额外{len(actual-expected)}张。')
        for _, r in a.iterrows():
            relative = Path(*Path(r.relative_path).parts[1:])
            source = roots[g] / relative
            st = source.stat()
            if st.st_size != r.size_bytes:
                raise ValueError(f'原图大小改变：{source}。不能沿用历史参考。')
            rows.append({'frame_id': r.frame_id, 'group': g, 'group_label': r.group_label,
                         'sequence_index': int(r.sequence_index), 'source_path': str(source),
                         'relative_path': r.relative_path, 'size_bytes': st.st_size, 'mtime_ns': st.st_mtime_ns})
    write_json(output / 'data/source_inventory_initial.json', rows)
    info.update(source_folders=normalized, profile=PROFILE, prepared=False)
    write_json(output / MARKER, info)
    run_script(output, 'inventory.py')
    fresh = pd.read_csv(output / 'data/图片清单.csv').set_index('frame_id').sort_index()
    expected = inv.set_index('frame_id').sort_index()
    for col in ['width', 'height', 'orientation', 'elapsed_s', 'interval_s']:
        pd.testing.assert_series_equal(fresh[col], expected[col], check_dtype=False)
    pd.testing.assert_series_equal(pd.to_datetime(fresh.timestamp), pd.to_datetime(expected.timestamp))
    samples = read_json(output / 'data/sample_selection.json')
    old_samples = read_json(CASE / 'data/sample_selection.json')
    assert {r['sample_id']: r['frame_id'] for r in samples} == {r['sample_id']: r['frame_id'] for r in old_samples}
    index = pd.read_csv(output / 'data/标注样本索引.csv')
    lookup = fresh.source_path.to_dict()
    index['source_path'] = index.frame_id.map(lookup)
    index['annotation_path'] = index.sample_id.map(lambda s: str(output / 'evidence' / f'{s}_计数标注.jpg'))
    index['unannotated_preview'] = ''
    index.to_csv(output / 'data/标注样本索引.csv', index=False, encoding='utf-8-sig')
    current_config = read_json(output / 'data/project_config.json')
    for group in current_config['groups']:
        group['source_folder'] = normalized[group['id']]
    current_config['replay_profile'] = PROFILE
    current_config['source_root'] = None
    write_json(output / 'data/project_config.json', current_config)
    review_doc = output / 'review/复核工具使用说明.md'
    review_text = review_doc.read_text(encoding='utf-8')
    review_text = review_text.replace(
        '双击结果目录的“复核照片.command”，会启动本机服务并打开浏览器。保留启动窗口，关闭服务后页面就不能继续读取原图。也可以在结果目录运行：',
        '使用Skill的review命令启动本机服务并打开浏览器。保留启动窗口，关闭服务后页面就不能继续读取原图。也可以在结果目录运行：')
    review_doc.write_text(review_text, encoding='utf-8')
    info.update(prepared=True, prepared_at=now())
    write_json(output / MARKER, info)
    check_runtime(output)
    print('原图清单与拍摄时间匹配，已准备复现项目：', output)
    return output


def check_runtime(output):
    manifest = read_json(CASE / 'detector_manifest.json')
    for rel, expected in manifest['files'].items():
        if digest(output / rel) != expected:
            raise ValueError(f'检测器已变化：{rel}。请建立新版本并重新核验。')
    for script in (FROZEN / 'scripts').glob('*.py'):
        if digest(output / 'scripts' / script.name) != digest(script):
            raise ValueError(f'冻结运行脚本已变化：{script.name}。严格复现不能继承修改前的验收状态。')


def smoke(output):
    verify_assets()
    output, info = project(output)
    if not info.get('prepared'):
        raise ValueError('先完成 prepare。')
    check_runtime(output)
    sys.path.insert(0, str(output / 'scripts'))
    import run_batch
    import pandas as pd
    assert run_batch.fingerprint()['fingerprint'] == FINGERPRINT
    inv = pd.read_csv(output / 'data/图片清单.csv').set_index('frame_id')
    refs = baseline().set_index('frame_id')
    selected = {'A_cal_01', 'A_val_15', 'B_cal_01', 'B_val_15', 'C_cal_01', 'C_val_29', 'C_val_30'}
    rows = [r for r in read_json(output / 'data/sample_selection.json') if r['sample_id'] in selected]
    assert len(rows) == 7
    outcomes = []
    with (output / 'qa/原图冒烟检查.jsonl').open('w', encoding='utf-8') as f:
        for s in rows:
            r = inv.loc[s['frame_id']].to_dict()
            r.update(frame_id=s['frame_id'], sample_id=s['sample_id'], use_cache=False)
            prediction = run_batch.process(r)
            f.write(json.dumps(prediction, ensure_ascii=False) + '\n')
            ref = refs.loc[r['frame_id']]
            assert prediction['status'] == 'processed', prediction.get('error')
            assert prediction['sha256'] == ref.source_sha256, r['frame_id']
            for key in ['auto_left_count', 'auto_right_count', 'auto_other_count']:
                assert prediction[key] == ref[key], (r['frame_id'], key, prediction[key], ref[key])
            outcomes.append({'sample_id': s['sample_id'], 'hash_and_counts_match': True})
            print('原图冒烟检查通过：', s['sample_id'], flush=True)
    write_json(output / 'qa/原图冒烟检查汇总.json', {'checked_at': now(), 'frames': outcomes,
        'scope': '冻结检测器数值复现；不是新的准确性验收'})


def verify_records(output, require_complete=True):
    check_runtime(output)
    refs = baseline().set_index('frame_id')
    path = output / 'data' / f'run_{FINGERPRINT}' / 'full_checkpoint.jsonl'
    if not path.exists():
        if require_complete:
            raise ValueError('尚无全量逐图记录。')
        return None
    seen = set(); mismatches = []
    with path.open(encoding='utf-8') as f:
        for line in f:
            try:
                r = json.loads(line)
            except json.JSONDecodeError:
                if not require_complete and not line.endswith('\n') and f.read() == '':
                    print('检测到中断尾行；原运行器将备份尾行后续跑。')
                    break
                raise ValueError('检查点存在损坏行。保留文件，先定位损坏，不能跳过后宣称完整。')
            fid = r['frame_id']
            if fid in seen or fid not in refs.index:
                raise ValueError(f'逐图记录重复或标识未知：{fid}')
            seen.add(fid); ref = refs.loc[fid]
            if r.get('status') != 'processed' or r.get('sha256') != ref.source_sha256:
                mismatches.append({'frame_id': fid, 'reason': '读取状态或原图SHA256不匹配'})
                continue
            if not r.get('source_unchanged'):
                mismatches.append({'frame_id': fid, 'reason': '原图大小或修改时间在准备后发生变化'})
            effective_quality = not r.get('quality_reasons') and not (r['group'] == 'C' and r.get('sharpness', 0) < 90)
            if effective_quality != bool(ref.quality_ok):
                mismatches.append({'frame_id': fid, 'reason': '有效画面筛选与冻结案例不同'})
            for key in ['auto_left_count', 'auto_right_count', 'auto_other_count', 'suspect_count']:
                if r.get(key) != ref[key]:
                    mismatches.append({'frame_id': fid, 'reason': key, 'actual': r.get(key), 'expected': int(ref[key])})
            ds = r.get('detections', [])
            for side in ['left', 'right']:
                assert sum(d.get('side') == side for d in ds) == r['auto_' + side + '_count']
    result = {'checked_at': now(), 'n_records': len(seen), 'expected_total': TOTAL,
              'all_records_present': len(seen) == TOTAL, 'mismatch_count': len(mismatches),
              'first_100_mismatches': mismatches[:100], 'accuracy_acceptance_is_separate': True}
    write_json(output / 'qa/原图复现一致性.json', result)
    if mismatches or (require_complete and len(seen) != TOTAL):
        raise ValueError('逐图复现不匹配或未完成，已保存差异；禁止继承历史验收结论或导出为已完成结果。')
    return result


def run_detection(output, workers):
    verify_assets()
    output, info = project(output)
    if not info.get('prepared'):
        raise ValueError('先完成 prepare。')
    if not 1 <= workers <= 32:
        raise ValueError('workers应为1到32；本次原运行使用16。')
    verify_records(output, require_complete=False)
    run_script(output, 'run_batch.py', '--mode', 'full', '--workers', workers)
    verify_records(output)


def export(output):
    verify_assets()
    output, _ = project(output)
    verify_records(output)
    run_script(output, 'test_statistics.py')
    run_script(output, 'export_results.py')
    run_script(output, 'verify_deliverables.py')
    compare_statistics(output)
    check_figures(output)
    report = output / '中文分析报告.md'
    text = report.read_text(encoding='utf-8')
    text = text.replace('生成日期：2026-09-16。', f'本次复现生成日期：{now()}；冻结方案日期：2026-09-16。')
    text = text.replace('## 1. 数据范围', '本次复用历史模型目视参考记录，没有重新进行独立人工专家核验。原始照片SHA256及逐图自动计数已与冻结案例逐一核对。\n\n## 1. 数据范围', 1)
    text = text.replace('`qa/`保留未标注参考视图、局部放大和验证记录。',
                        '`qa/`保留本次复现的验证记录；历史未标注参考视图和局部放大未在此重新生成。')
    report.write_text(text, encoding='utf-8')
    instructions = output / '复现说明.md'
    instructions.write_text(
        '# 本地复现\n\n使用本次Skill的 `scripts/workflow.py` 入口，依次prepare、smoke、run、export。'
        '该入口核对原图SHA256、冻结运行器、全量计数及统计基准。\n\n'
        f'Skill位置：{SKILL}\n\n完整命令见Skill的 references/reproduction.md。'
        '本次重用历史模型目视参考，不是新的独立专家验收；B/C仍未通过。\n', encoding='utf-8')
    write_json(output / 'qa/复现完成状态.json', {'completed_at': now(), 'technical_replay_complete': True,
               'all_groups_count_acceptance_completed': False, 'human_expert_acceptance_completed': False,
               'primary_group': 'A，沿用历史模型目视初检状态；B/C仍未通过'})
    refresh_delivery_manifest(output)
    print('原图复现、表格和图形导出完成；B/C仍未通过计数验收：', output)


def init_new(config_path, output):
    """Create only a new-study specification; never import historical approvals."""
    config = read_json(config_path)
    if config.get('profile') != 'new-study':
        raise ValueError('新实验配置必须为 new-study。')
    groups = config.get('groups', [])
    if not groups or len({g['id'] for g in groups}) != len(groups):
        raise ValueError('新实验分组为空或编号重复。')
    output = create_output(output, 'new-study', [g['source_folder'] for g in groups])
    config['validation_status'] = 'pending'
    config['historical_acceptance_inherited'] = False
    write_json(output / 'data/new_study_config.json', config)
    with (output / 'data/目视参考模板.csv').open('w', encoding='utf-8-sig', newline='') as f:
        csv.writer(f).writerow(['sample_id', 'frame_id', 'source_path', 'group', 'split', 'left_count',
                               'right_count', 'other_count', 'uncertain_count', 'reviewer', 'review_time',
                               'viewed_predictions_before_reference', 'countable', 'reason', 'coordinates_file'])
    (output / '新实验待办.md').write_text(
        '# 新实验\n\n这只是配置与参考表模板，尚未清点、检测或验收。'
        '不能运行冻结A/B/C检测器来继承旧结论。按Skill的新数据流程检查原图、场景边界、时间，'
        '建立不重叠的校准与留出集，逐体标注后适配检测器，再进行新的留出验收。\n', encoding='utf-8')
    print('新实验模板已创建，尚未计数或验收：', output)


def main():
    p = argparse.ArgumentParser(description=__doc__)
    sub = p.add_subparsers(dest='command', required=True)
    sub.add_parser('doctor', help='检查本地环境及冻结资源，不安装依赖')
    a = sub.add_parser('replot', help='不读取原图，复绘冻结的全量逐图计数'); a.add_argument('--output', required=True)
    a = sub.add_parser('prepare', help='在新目录准备严格的原图复现'); a.add_argument('--config', default=str(SKILL/'assets/fafu-2023.example.json')); a.add_argument('--output', required=True)
    for name in ['smoke', 'run', 'export', 'verify', 'review']:
        a = sub.add_parser(name); a.add_argument('--workspace', required=True)
        if name == 'run': a.add_argument('--workers', type=int, default=8)
        if name == 'review': a.add_argument('--port', type=int, default=8768)
    a = sub.add_parser('init-new'); a.add_argument('--config', required=True); a.add_argument('--output', required=True)
    args = p.parse_args()
    try:
        if args.command == 'doctor': doctor()
        elif args.command == 'replot': replot(args.output)
        elif args.command == 'prepare': prepare(args.config, args.output)
        elif args.command == 'smoke': smoke(args.workspace)
        elif args.command == 'run': run_detection(args.workspace, args.workers)
        elif args.command == 'export': export(args.workspace)
        elif args.command == 'verify':
            verify_assets(); output, _ = project(args.workspace); print(json.dumps(verify_records(output), ensure_ascii=False, indent=2))
        elif args.command == 'review':
            output, _ = project(args.workspace)
            if not (output/'data/自动检测索引.sqlite').exists(): raise ValueError('先完成全量导出。')
            run_script(output, 'review_server.py', '--port', args.port, '--open-browser')
        elif args.command == 'init-new': init_new(args.config, args.output)
    except (ValueError, AssertionError, OSError, subprocess.CalledProcessError) as exc:
        print(f'未完成：{exc}', file=sys.stderr)
        sys.exit(1)


if __name__ == '__main__':
    main()
